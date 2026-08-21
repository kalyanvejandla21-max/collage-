import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_btech_presentation():
    prs = Presentation()
    # 16:9 Widescreen aspect ratio (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    BG_DARK = RGBColor(11, 15, 25)          # Deep Midnight Navy #0B0F19
    CARD_BG = RGBColor(18, 26, 43)          # Slate Card #121A2B
    CARD_BORDER = RGBColor(37, 99, 235)     # Royal Blue Border #2563EB
    COLOR_CYAN = RGBColor(6, 182, 212)      # Cyber Cyan #06B6D4
    COLOR_GOLD = RGBColor(245, 158, 11)     # Amber Gold #F59E0B
    COLOR_PURPLE = RGBColor(139, 92, 246)   # Violet Purple #8B5CF6
    COLOR_EMERALD = RGBColor(16, 185, 129)  # Emerald Green #10B981
    COLOR_RED = RGBColor(239, 68, 68)       # Warning Red #EF4444
    TEXT_WHITE = RGBColor(248, 250, 252)    # White #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184)    # Slate Muted #94A3B8

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_slide_header(slide, title_text, slide_num_str):
        # Top Accent Line
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
        accent.fill.solid()
        accent.fill.fore_color.rgb = COLOR_CYAN
        accent.line.fill.background()

        # Header Box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.733), Inches(1.0))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = f"BATCH 14  •  B.TECH PROJECT REVIEW  •  {slide_num_str}".upper()
        p0.font.name = 'Montserrat'
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_GOLD

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.name = 'Montserrat'
        p1.font.size = Pt(28)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.35))
        ftf = footer_box.text_frame
        ftf.word_wrap = True
        fp = ftf.paragraphs[0]
        fp.text = "Batch No. 14 | Academic Year 2026–2027 | Guide: Ms. K. Neeharika"
        fp.font.name = 'Inter'
        fp.font.size = Pt(10)
        fp.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, title, items, highlight_color=COLOR_CYAN):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = highlight_color
        card.line.width = Pt(1.5)

        box = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.25), Inches(width - 0.5), Inches(height - 0.5))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        if title:
            p_title = tf.paragraphs[0]
            p_title.text = title
            p_title.font.name = 'Montserrat'
            p_title.font.size = Pt(16)
            p_title.font.bold = True
            p_title.font.color.rgb = highlight_color
            p_title.space_after = Pt(10)

        first_item = True if title else False
        for item in items:
            p_item = tf.paragraphs[0] if (not title and first_item) else tf.add_paragraph()
            first_item = False
            p_item.font.name = 'Inter'
            p_item.font.size = Pt(13)
            p_item.space_after = Pt(6)

            if isinstance(item, tuple):
                heading, body = item
                r1 = p_item.add_run()
                r1.text = "• " + heading + ": "
                r1.font.bold = True
                r1.font.color.rgb = TEXT_WHITE

                r2 = p_item.add_run()
                r2.text = body
                r2.font.color.rgb = TEXT_MUTED
            else:
                r = p_item.add_run()
                r.text = "• " + str(item)
                r.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE 1 — TITLE / PROJECT INTRODUCTION
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Decorative background circles/accents
    decor1 = slide1.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.5), Inches(-1.5), Inches(6.5), Inches(6.5))
    decor1.fill.solid()
    decor1.fill.fore_color.rgb = RGBColor(15, 28, 60)
    decor1.line.fill.background()

    decor2 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.3), Inches(13.333), Inches(0.2))
    decor2.fill.solid()
    decor2.fill.fore_color.rgb = COLOR_CYAN
    decor2.line.fill.background()

    # Left Column: Project Title & Details
    t_box = slide1.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(7.5), Inches(5.8))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "B.TECH FINAL YEAR PROJECT PRESENTATION"
    p0.font.name = 'Montserrat'
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_GOLD
    p0.space_after = Pt(10)

    p1 = tf1.add_paragraph()
    p1.text = "WEB-BASED ONLINE EXAMINATION PORTAL"
    p1.font.name = 'Montserrat'
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(8)

    p2 = tf1.add_paragraph()
    p2.text = "Randomized Question Generation and Exam Security"
    p2.font.name = 'Montserrat'
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_CYAN
    p2.space_after = Pt(24)

    # Metadata Block
    p_meta = tf1.add_paragraph()
    r_meta1 = p_meta.add_run()
    r_meta1.text = "PROJECT BATCH NO: "
    r_meta1.font.bold = True
    r_meta1.font.size = Pt(13)
    r_meta1.font.color.rgb = COLOR_GOLD

    r_meta2 = p_meta.add_run()
    r_meta2.text = "Batch 14    |    "
    r_meta2.font.size = Pt(13)
    r_meta2.font.color.rgb = TEXT_WHITE

    r_meta3 = p_meta.add_run()
    r_meta3.text = "ACADEMIC YEAR: "
    r_meta3.font.bold = True
    r_meta3.font.size = Pt(13)
    r_meta3.font.color.rgb = COLOR_GOLD

    r_meta4 = p_meta.add_run()
    r_meta4.text = "2026–2027\n"
    r_meta4.font.size = Pt(13)
    r_meta4.font.color.rgb = TEXT_WHITE

    p_guide = tf1.add_paragraph()
    r_g1 = p_guide.add_run()
    r_g1.text = "PROJECT GUIDE: "
    r_g1.font.bold = True
    r_g1.font.size = Pt(13)
    r_g1.font.color.rgb = COLOR_GOLD

    r_g2 = p_guide.add_run()
    r_g2.text = "Ms. K. Neeharika"
    r_g2.font.size = Pt(13)
    r_g2.font.color.rgb = TEXT_WHITE
    p_guide.space_after = Pt(20)

    # Team Members Card on Slide 1 Right
    add_card(slide1, 8.5, 1.2, 4.0, 5.2, "TEAM MEMBERS", [
        ("G. Uday Kiran", "24HP1A0565"),
        ("K. Nagarjuna", "24HP1A0552"),
        ("Ranga Swamy", "24HP1A0561"),
        ("V. Kalyan", "24HP1A0541")
    ], COLOR_PURPLE)

    # Visual Badges on Bottom Left of Slide 1
    badge_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.7), Inches(7.3), Inches(1.0))
    badge_box.fill.solid()
    badge_box.fill.fore_color.rgb = CARD_BG
    badge_box.line.color.rgb = CARD_BORDER
    btf = badge_box.text_frame
    btf.word_wrap = True
    bp = btf.paragraphs[0]
    bp.alignment = PP_ALIGN.CENTER
    bp.text = "💻 Online Exam Dashboard  |  🛡 Fullscreen & Tab Monitoring  |  ⏱ 30-Min Timer  |  🔀 Fisher-Yates MCQ Shuffle"
    bp.font.name = 'Inter'
    bp.font.size = Pt(11)
    bp.font.bold = True
    bp.font.color.rgb = COLOR_CYAN

    # =========================================================================
    # SLIDE 2 — PROBLEM STATEMENT
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_slide_header(slide2, "PROBLEM STATEMENT", "SLIDE 02")

    # Problems List (Left Box)
    add_card(slide2, 0.8, 1.5, 6.0, 4.1, "Conventional Examination Drawbacks", [
        ("Time-Consuming Setup", "Manual question paper preparation requires excessive faculty effort."),
        ("Predictable Question Paper", "Students receive identical or static question sequences."),
        ("Reduced Assessment Value", "Repeated questions compromise exam integrity and evaluation effectiveness."),
        ("Manual Evaluation Delay", "Physical grading introduces delays, human errors, and resource overhead."),
        ("Limited Security Monitoring", "Basic systems lack active tab-switching and window-blur detection."),
        ("Delayed Results", "Students wait days/weeks for score publication.")
    ], COLOR_RED)

    # Problem -> Challenges -> Need Diagram (Right Box)
    add_card(slide2, 7.1, 1.5, 5.4, 4.1, "Problem → Challenges → Need", [
        ("1. Manual Work", "Paper creation & manual grading burden faculty."),
        ("2. Repetition", "Static question papers facilitate unfair practices."),
        ("3. Security Issues", "Tab switching & window leaving undetected."),
        ("4. Delayed Results", "No real-time feedback or automated scoring.")
    ], COLOR_GOLD)

    # Main Message Banner (Bottom Container)
    msg_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.75), Inches(11.733), Inches(1.0))
    msg_box.fill.solid()
    msg_box.fill.fore_color.rgb = RGBColor(24, 34, 56)
    msg_box.line.color.rgb = COLOR_CYAN
    msg_box.line.width = Pt(1.5)
    mtf = msg_box.text_frame
    mtf.word_wrap = True
    mp0 = mtf.paragraphs[0]
    mp0.text = "CORE MISSION REQUIREMENT:"
    mp0.font.name = 'Montserrat'
    mp0.font.size = Pt(10)
    mp0.font.bold = True
    mp0.font.color.rgb = COLOR_GOLD

    mp1 = mtf.add_paragraph()
    mp1.text = "\"There is a need for a centralized, automated and secure online examination platform that can generate randomized question papers, monitor examination behavior and evaluate answers instantly.\""
    mp1.font.name = 'Inter'
    mp1.font.size = Pt(13)
    mp1.font.bold = True
    mp1.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE 3 — PROPOSED SYSTEM & WORKFLOW
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_slide_header(slide3, "PROPOSED SYSTEM & WORKFLOW", "SLIDE 03")

    # Overview Box
    add_card(slide3, 0.8, 1.5, 11.733, 1.0, "Centralized Web-Based Examination Platform", [
        "Provides an end-to-end automated platform integrating student authentication, randomized MCQ extraction, live browser-level security enforcement, automated evaluation, and instant result publishing."
    ], COLOR_CYAN)

    # 8-Node Flowchart (2 rows of 4 nodes)
    flow_steps = [
        ("1. LOGIN", "Credential validation", COLOR_CYAN),
        ("2. SUBJECT SELECT", "Pick exam module", COLOR_CYAN),
        ("3. QUESTION BANK", "Load MCQ database", COLOR_CYAN),
        ("4. RANDOM 20 MCQs", "Extract question subset", COLOR_PURPLE),
        ("5. SHUFFLE MCQs", "Fisher-Yates shuffle", COLOR_PURPLE),
        ("6. SECURE EXAM", "Fullscreen & focus lock", COLOR_GOLD),
        ("7. AUTO EVALUATION", "Instant score compute", COLOR_EMERALD),
        ("8. RESULT", "Immediate feedback", COLOR_EMERALD)
    ]

    node_w = Inches(2.7)
    node_h = Inches(1.8)

    # Row 1 (Nodes 1-4)
    for i in range(4):
        left_pos = 0.8 + i * 3.0
        title, sub, color = flow_steps[i]

        box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_pos), Inches(2.7), node_w, node_h)
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = color
        box.line.width = Pt(1.5)

        tf = box.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        p1.text = title
        p1.font.name = 'Montserrat'
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = color
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.text = sub
        p2.font.name = 'Inter'
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED

    # Row 2 (Nodes 5-8)
    for i in range(4):
        left_pos = 0.8 + i * 3.0
        title, sub, color = flow_steps[i + 4]

        box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_pos), Inches(4.8), node_w, node_h)
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = color
        box.line.width = Pt(1.5)

        tf = box.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        p1.text = title
        p1.font.name = 'Montserrat'
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = color
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.text = sub
        p2.font.name = 'Inter'
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 4 — RANDOMIZED MCQ SYSTEM & KEY FEATURES
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_slide_header(slide4, "RANDOMIZED MCQ SYSTEM & KEY FEATURES", "SLIDE 04")

    # Left Column: Fisher-Yates & Randomization Rules
    add_card(slide4, 0.8, 1.5, 5.7, 5.2, "Fisher-Yates Shuffle & Randomization", [
        ("Subject Question Banks", "Maintains MCQ banks for subjects (Computer Networks, Quantum Computing)."),
        ("Exact 20-Question Extraction", "Randomly selects exactly 20 non-duplicate questions per test attempt."),
        ("Question Sequence Shuffle", "Fisher-Yates algorithm randomizes question ordering."),
        ("Option Shuffling (A, B, C, D)", "All options shuffled independently per question."),
        ("Correct Answer Mapping", "Preserves underlying correct answer index after option shuffle."),
        ("Unique Test Instance", "Every student attempt produces a unique, un-copyable paper.")
    ], COLOR_PURPLE)

    # Right Column: Examination Engine Features
    add_card(slide4, 6.8, 1.5, 5.7, 5.2, "Core Examination Engine Features", [
        ("30-Minute Countdown Timer", "Live countdown with automated force submission on expiry."),
        ("Question Navigation Palette", "Color-coded grid for Answered / Unanswered tracking."),
        ("Previous / Next Controls", "Seamless switching between questions with answer state persistence."),
        ("Auto-Save Engine", "Student responses saved instantly in background local/session storage."),
        ("Review Unanswered Items", "Quick filter to jump to skipped questions before submitting."),
        ("Instant Score Calculation", "Immediate evaluation upon student submit click.")
    ], COLOR_CYAN)

    # =========================================================================
    # SLIDE 5 — EXAM SECURITY & MONITORING
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_slide_header(slide5, "EXAM SECURITY & MONITORING", "SLIDE 05")

    # Left Box: Security Mechanisms & Web APIs
    add_card(slide5, 0.8, 1.5, 6.0, 5.2, "Browser Security Mechanisms & Web APIs", [
        ("1. Fullscreen Enforcement", "Browser requests fullscreen. Exit triggers detection, warning prompt, & log entry."),
        ("2. Tab Switch & Focus Detection", "Detects window blur, tab switching, and browser losing active focus."),
        ("Web APIs Employed", "Fullscreen API (fullscreenchange), Page Visibility API (visibilitychange), Window Blur/Focus events."),
        ("3-Tier Violation Protocol", "1st: Warning → 2nd: Strong Warning → 3rd: Final Warning."),
        ("Security Audit Log", "Logs violation counter, event type, exact timestamp, and student ID."),
        ("Scope Boundary Note", "Provides robust browser-level monitoring; does not monitor external physical hardware.")
    ], COLOR_GOLD)

    # Right Box: Security Dashboard Mockup Card
    dash = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(1.5), Inches(5.4), Inches(5.2))
    dash.fill.solid()
    dash.fill.fore_color.rgb = CARD_BG
    dash.line.color.rgb = COLOR_RED
    dash.line.width = Pt(2.0)

    dtf = dash.text_frame
    dtf.word_wrap = True

    dp0 = dtf.paragraphs[0]
    dp0.text = "🔒 LIVE EXAM SECURITY MONITORING DASHBOARD"
    dp0.font.name = 'Montserrat'
    dp0.font.size = Pt(14)
    dp0.font.bold = True
    dp0.font.color.rgb = COLOR_RED
    dp0.space_after = Pt(14)

    status_items = [
        ("FULLSCREEN STATUS", "🟢 ACTIVE (LOCK ENGAGED)", COLOR_EMERALD),
        ("PAGE VISIBILITY", "🟢 FOCUS DETECTED (TAB ACTIVE)", COLOR_EMERALD),
        ("VIOLATION COUNTER", "⚠️ 1 / 3 WARNINGS ISSUED", COLOR_GOLD),
        ("SECURITY EVENT LOG", "• 09:14:22 AM - FULLSCREEN_EXIT DETECTED\n• 09:14:25 AM - FULLSCREEN_RESTORED", COLOR_CYAN),
        ("MONITORING APIS", "Fullscreen API | Page Visibility API | Focus/Blur", TEXT_MUTED)
    ]

    for label, val, val_color in status_items:
        p_lbl = dtf.add_paragraph()
        r_lbl = p_lbl.add_run()
        r_lbl.text = label + "\n"
        r_lbl.font.name = 'Montserrat'
        r_lbl.font.size = Pt(11)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = COLOR_GOLD

        r_val = p_lbl.add_run()
        r_val.text = val + "\n"
        r_val.font.name = 'Inter'
        r_val.font.size = Pt(12)
        r_val.font.bold = True
        r_val.font.color.rgb = val_color
        p_lbl.space_after = Pt(8)

    # =========================================================================
    # SLIDE 6 — TECHNOLOGY STACK & SYSTEM ARCHITECTURE
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_slide_header(slide6, "TECHNOLOGY STACK & SYSTEM ARCHITECTURE", "SLIDE 06")

    # 4 Tech Cards (Top Half)
    add_card(slide6, 0.8, 1.5, 2.7, 3.2, "FRONTEND", [
        ("Core Tech", "HTML5, CSS3, JavaScript, React.js"),
        ("Usage", "User Interface, Login, Subject selection, Exam screen, Timer, Result view")
    ], COLOR_CYAN)

    add_card(slide6, 3.8, 1.5, 2.7, 3.2, "BACKEND", [
        ("Core Tech", "Node.js, Express.js"),
        ("Usage", "REST APIs, Exam session management, Question retrieval, Scoring logic")
    ], COLOR_PURPLE)

    add_card(slide6, 6.8, 1.5, 2.7, 3.2, "DATABASE", [
        ("Core Tech", "MySQL or MongoDB"),
        ("Usage", "Student info, Question bank, Sessions, Answers, Security logs")
    ], COLOR_GOLD)

    add_card(slide6, 9.8, 1.5, 2.7, 3.2, "DATA & SECURITY", [
        ("Files", "Excel-based MCQ files (Question, Opt A-D, Correct Index)"),
        ("APIs", "Fullscreen & Page Visibility APIs")
    ], COLOR_EMERALD)

    # Layered Architecture Diagram (Bottom Half)
    arch_box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.9), Inches(11.733), Inches(1.8))
    arch_box.fill.solid()
    arch_box.fill.fore_color.rgb = CARD_BG
    arch_box.line.color.rgb = CARD_BORDER
    arch_box.line.width = Pt(1.5)

    atf = arch_box.text_frame
    atf.word_wrap = True
    ap0 = atf.paragraphs[0]
    ap0.alignment = PP_ALIGN.CENTER
    ap0.text = "LAYERED SYSTEM ARCHITECTURE DIAGRAM"
    ap0.font.name = 'Montserrat'
    ap0.font.size = Pt(12)
    ap0.font.bold = True
    ap0.font.color.rgb = COLOR_GOLD
    ap0.space_after = Pt(10)

    ap1 = atf.add_paragraph()
    ap1.alignment = PP_ALIGN.CENTER
    ap1.text = "STUDENT  ──▶  REACT FRONTEND (Security APIs Enabled)  ──▶  NODE.JS / EXPRESS REST API  ──▶  DATABASE / EXCEL QUESTION BANK"
    ap1.font.name = 'Montserrat'
    ap1.font.size = Pt(13)
    ap1.font.bold = True
    ap1.font.color.rgb = COLOR_CYAN

    # =========================================================================
    # SLIDE 7 — RESULTS, BENEFITS & FUTURE SCOPE
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_slide_header(slide7, "RESULTS, BENEFITS & FUTURE SCOPE", "SLIDE 07")

    # Column 1: Achieved Benefits & Results Card Breakdown
    add_card(slide7, 0.8, 1.5, 3.6, 4.0, "Achieved Benefits", [
        ("Paper Randomization", "Un-copyable question sequences."),
        ("Zero Evaluation Delay", "Instant result generation."),
        ("Faculty Workload", "Eliminates manual paper grading."),
        ("Browser Security", "Tracks tab switches & fullscreen exit.")
    ], COLOR_EMERALD)

    # Column 2: Result Page Output Breakdown
    add_card(slide7, 4.6, 1.5, 3.8, 4.0, "Result Screen Output", [
        ("Student Details", "Student Name, Reg Number, Subject"),
        ("Exam Metrics", "Total Questions (20), Correct & Wrong Count"),
        ("Score Computation", "Marks Obtained, Percentage, Pass/Fail")
    ], COLOR_CYAN)

    # Column 3: Future Scope
    add_card(slide7, 8.6, 1.5, 3.9, 4.0, "Future Scope", [
        ("Admin Dashboard", "Faculty question bank & analytics management"),
        ("Enhanced Security", "Biometric student authentication & AI proctoring"),
        ("Scaling", "Cloud deployment on AWS & multi-subject expansion")
    ], COLOR_PURPLE)

    # Thank You Banner at Bottom
    thanks_box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.65), Inches(11.733), Inches(1.2))
    thanks_box.fill.solid()
    thanks_box.fill.fore_color.rgb = RGBColor(24, 34, 56)
    thanks_box.line.color.rgb = COLOR_GOLD
    thanks_box.line.width = Pt(1.5)

    ttf = thanks_box.text_frame
    ttf.word_wrap = True

    tp0 = ttf.paragraphs[0]
    tp0.alignment = PP_ALIGN.CENTER
    tp0.text = "THANK YOU!"
    tp0.font.name = 'Montserrat'
    tp0.font.size = Pt(20)
    tp0.font.bold = True
    tp0.font.color.rgb = COLOR_GOLD
    tp0.space_after = Pt(2)

    tp1 = ttf.add_paragraph()
    tp1.alignment = PP_ALIGN.CENTER
    tp1.text = "Team: G. Uday Kiran | K. Nagarjuna | Ranga Swamy | V. Kalyan    •    Batch 14 | Academic Year 2026–2027 | Guide: Ms. K. Neeharika"
    tp1.font.name = 'Inter'
    tp1.font.size = Pt(11)
    tp1.font.bold = True
    tp1.font.color.rgb = TEXT_WHITE

    output_path = "Batch14_BTech_Project_Presentation.pptx"
    prs.save(output_path)
    print(f"B.Tech presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_btech_presentation()
