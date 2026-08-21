import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # Set 16:9 widescreen aspect ratio (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    COLOR_BG = RGBColor(15, 23, 42)        # Slate 900
    COLOR_CARD = RGBColor(30, 41, 59)      # Slate 800
    COLOR_CARD_BORDER = RGBColor(51, 65, 85) # Slate 700
    COLOR_GOLD = RGBColor(245, 158, 11)   # Amber 500
    COLOR_CYAN = RGBColor(6, 182, 212)    # Cyan 500
    COLOR_PURPLE = RGBColor(139, 92, 246) # Violet 500
    COLOR_TEXT_MAIN = RGBColor(248, 250, 252) # Slate 50
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184) # Slate 400
    COLOR_WHITE = RGBColor(255, 255, 255)

    blank_layout = prs.slide_layouts[6] # Blank slide layout

    def add_background_and_header(slide, title_text, category_text="ALIET EXAM PORTAL"):
        # Dark Background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()

        # Top Accent Line
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.1))
        accent.fill.solid()
        accent.fill.fore_color.rgb = COLOR_GOLD
        accent.line.fill.background()

        # Header Box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.0))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Category
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.name = 'Calibri'
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_GOLD

        # Title
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.name = 'Calibri'
        p1.font.size = Pt(26)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_MAIN

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.4))
        ftf = footer_box.text_frame
        ftf.word_wrap = True
        fp = ftf.paragraphs[0]
        fp.text = "Andhra Loyola Institute of Engineering and Technology | Examination & Assessment System"
        fp.font.name = 'Calibri'
        fp.font.size = Pt(10)
        fp.font.color.rgb = COLOR_TEXT_MUTED

    def add_card(slide, left, top, width, height, title, items, highlight_color=COLOR_CYAN):
        # Card Background
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_CARD_BORDER
        card.line.width = Pt(1.5)

        # Content Box
        box = slide.shapes.add_textbox(Inches(left + 0.3), Inches(top + 0.3), Inches(width - 0.6), Inches(height - 0.6))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Card Title
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.name = 'Calibri'
        p_title.font.size = Pt(18)
        p_title.font.bold = True
        p_title.font.color.rgb = highlight_color
        p_title.space_after = Pt(14)

        # Card Items
        for item in items:
            p_item = tf.add_paragraph()
            p_item.font.name = 'Calibri'
            p_item.font.size = Pt(13)
            p_item.space_after = Pt(8)

            if isinstance(item, tuple):
                heading, body = item
                run1 = p_item.add_run()
                run1.text = "• " + heading + ": "
                run1.font.bold = True
                run1.font.color.rgb = COLOR_TEXT_MAIN

                run2 = p_item.add_run()
                run2.text = body
                run2.font.color.rgb = COLOR_TEXT_MUTED
            else:
                run = p_item.add_run()
                run.text = "• " + str(item)
                run.font.color.rgb = COLOR_TEXT_MAIN

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_BG
    bg1.line.fill.background()

    # Decorative glow accent shapes
    accent1 = slide1.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9), Inches(-1), Inches(6), Inches(6))
    accent1.fill.solid()
    accent1.fill.fore_color.rgb = RGBColor(30, 58, 138)
    accent1.line.fill.background()

    accent2 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.3), Inches(13.333), Inches(0.2))
    accent2.fill.solid()
    accent2.fill.fore_color.rgb = COLOR_GOLD
    accent2.line.fill.background()

    # Title Text Frame
    t_box = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(4.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "ANDHRA LOYOLA INSTITUTE OF ENGINEERING AND TECHNOLOGY"
    p0.font.name = 'Calibri'
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_GOLD
    p0.space_after = Pt(12)

    p1 = tf1.add_paragraph()
    p1.text = "ALIET Online Examination Portal"
    p1.font.name = 'Calibri'
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.space_after = Pt(12)

    p2 = tf1.add_paragraph()
    p2.text = "Comprehensive Next-Generation Computer-Based Testing & Analytics Platform"
    p2.font.name = 'Calibri'
    p2.font.size = Pt(20)
    p2.font.color.rgb = COLOR_CYAN
    p2.space_after = Pt(28)

    p3 = tf1.add_paragraph()
    p3.text = "Subjects Covered: Computer Networks | Quantum Computing | Finite Automata"
    p3.font.name = 'Calibri'
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 2: Core System Architecture
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_background_and_header(slide2, "1. Core System Architecture & Technology Stack")

    add_card(slide2, 0.8, 1.6, 5.6, 4.8, "Frontend & User Interface", [
        ("Technology", "HTML5, Vanilla CSS3, Modern JavaScript ES6+"),
        ("UI Aesthetics", "Dark mode glassmorphism, animated glow effects"),
        ("Interactivity", "Particle explosion graphics, self-drawing neon SVG icons"),
        ("Responsive Layout", "Grid & Flexbox containers with zero latency rendering")
    ], COLOR_GOLD)

    add_card(slide2, 6.8, 1.6, 5.7, 4.8, "Backend, Database & Analytics", [
        ("Server Stack", "Node.js with Express.js REST API layer"),
        ("Persistence", "MongoDB with Mongoose schemas for Exam Results"),
        ("Data Export", "SheetJS (xlsx) workbook engine for live reports"),
        ("Security", "CORS enabled, sanitized input validation & rate limiting")
    ], COLOR_CYAN)

    # ==========================================
    # SLIDE 3: Multi-Subject Question Bank
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_background_and_header(slide3, "2. Comprehensive Multi-Subject Question Bank")

    add_card(slide3, 0.8, 1.6, 3.6, 4.8, "Computer Networks", [
        ("Scope", "50 Comprehensive MCQs"),
        ("Topics", "OSI Model layers, Topologies (Star, Bus, Ring), TCP/IP"),
        ("Media Types", "Twisted pair, Coaxial, Fiber optics bandwidth"),
        ("Routing", "Network layer protocols & IP addressing")
    ], COLOR_CYAN)

    add_card(slide3, 4.8, 1.6, 3.6, 4.8, "Quantum Computing", [
        ("Scope", "50 Specialized MCQs"),
        ("Topics", "Qubits, Superposition, Entanglement"),
        ("Mathematical Basis", "Unitary operators, Matrices, Eigenvalues"),
        ("Interdisciplinary", "Physics, Math, CS & Central Dogma links")
    ], COLOR_PURPLE)

    add_card(slide3, 8.8, 1.6, 3.7, 4.8, "Finite Automata", [
        ("Scope", "Curated Theory Questions"),
        ("Topics", "DFA vs NFA, Epsilon transitions, 5-tuple models"),
        ("Languages", "Regular Languages & Context-Free Grammars"),
        ("Machines", "Pushdown Automata & Turing Machine scope")
    ], COLOR_GOLD)

    # ==========================================
    # SLIDE 4: Interactive Exam Portal Experience
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_background_and_header(slide4, "3. Interactive Exam Portal & Test Experience")

    add_card(slide4, 0.8, 1.6, 5.6, 4.8, "Student Examination Features", [
        ("Cinematic Intro", "Self-drawing grad cap SVG and canvas particle burst"),
        ("Subject Selector", "Dynamic tab switching for instant question loading"),
        ("Exam Timer", "Live countdown timer with automatic force submission"),
        ("Question Palette", "Color-coded indicators for Answered / Unanswered items")
    ], COLOR_CYAN)

    add_card(slide4, 6.8, 1.6, 5.7, 4.8, "Test Execution & Integrity", [
        ("State Persistence", "Real-time saving of student option selections"),
        ("Option Shuffling", "Randomized answer choices per student session"),
        ("Instant Submit", "Modal confirmation prompt preventing accidental submission"),
        ("Cross-Browser", "Full compatibility across Mobile, Tablet, and Desktop")
    ], COLOR_PURPLE)

    # ==========================================
    # SLIDE 5: Automated Evaluation & Analytics
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_background_and_header(slide5, "4. Automated Evaluation & Detailed Analytics Engine")

    add_card(slide5, 0.8, 1.6, 5.6, 4.8, "Scoring & Performance Metrics", [
        ("Instant Grading", "Immediate evaluation upon test completion"),
        ("Metrics Calculated", "Total Score, Correct / Incorrect counts, Accuracy %"),
        ("Pass / Fail Gate", "Configurable percentage threshold calculation"),
        ("Visual Feedback", "Dynamic score rings and percentage progress indicators")
    ], COLOR_GOLD)

    add_card(slide5, 6.8, 1.6, 5.7, 4.8, "Detailed Answer Review Mode", [
        ("Question Breakdown", "Item-by-item analysis showing student pick vs correct choice"),
        ("Concept Explanations", "In-depth rationale provided for every question"),
        ("Learning Reinforcement", "Highlights mistake areas for targeted student revision"),
        ("Attempt History", "Stores historical attempts for student progress tracking")
    ], COLOR_CYAN)

    # ==========================================
    # SLIDE 6: Excel Synchronization & Export
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_background_and_header(slide6, "5. Excel Synchronization & Report Generation")

    add_card(slide6, 0.8, 1.6, 5.6, 4.8, "Database & Spreadsheet Integration", [
        ("SheetJS Engine", "In-browser spreadsheet compilation without external tools"),
        ("Generated Workbooks", "Exam_Results_Database.xlsx & Computer_Networks_50_MCQs.xlsx"),
        ("Data Schema", "Student ID, Name, Subject, Score, Accuracy %, Timestamp"),
        ("API Sync", "POST route /api/results syncs browser state directly to MongoDB")
    ], COLOR_PURPLE)

    add_card(slide6, 6.8, 1.6, 5.7, 4.8, "Administrator Capabilities", [
        ("Bulk Downloads", "One-click export of complete student cohort results"),
        ("Audit Trails", "Timestamps and attempt metadata logged per submission"),
        ("Offline Storage", "Excel files formatted for instant Excel / Google Sheets opening"),
        ("Automated Seeding", "seed.js script populates initial question databases")
    ], COLOR_GOLD)

    # ==========================================
    # SLIDE 7: Strategic Benefits & Future Roadmap
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_background_and_header(slide7, "6. Strategic Advantages & Future Roadmap")

    add_card(slide7, 0.8, 1.6, 5.6, 4.8, "Key Strategic Advantages", [
        ("Zero Overhead", "Eliminates physical paper evaluation and manual grading"),
        ("Scalable Platform", "Node/Express backend handles thousands of concurrent tests"),
        ("High Engagement", "Gamified cinematic visuals boost student satisfaction"),
        ("Modular Design", "Easily expandable to support new engineering subjects")
    ], COLOR_CYAN)

    add_card(slide7, 6.8, 1.6, 5.7, 4.8, "Future Roadmap", [
        ("AI Question Generator", "LLM-assisted question bank expansion and difficulty scaling"),
        ("Proctoring Guard", "Webcam facial verification and tab-switch anti-cheat monitoring"),
        ("Mobile App", "React Native cross-platform mobile app deployment"),
        ("Cloud Infrastructure", "Dockerized containerization on AWS / GCP for auto-scaling")
    ], COLOR_PURPLE)

    output_path = "ALIET_Examination_Portal_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_presentation()
